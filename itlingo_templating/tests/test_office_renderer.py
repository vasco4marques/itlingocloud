import io
import unittest

from docx import Document
from openpyxl import Workbook, load_workbook

from odoo.addons.itlingo_templating.services.docx_renderer import render_docx
from odoo.addons.itlingo_templating.services.xlsx_renderer import render_xlsx


class TestOfficeRenderers(unittest.TestCase):
    def test_docx_unknown_tag_reports_template_text(self):
        document = Document()
        document.add_paragraph("{%p fro x in y %}")
        template = io.BytesIO()
        document.save(template)

        with self.assertRaises(ValueError) as raised:
            render_docx(template.getvalue(), {})

        message = str(raised.exception)
        self.assertTrue(message.startswith("DOCX template error:"), message)
        self.assertIn("{%p fro x in y %}", message)
    def test_docx_unknown_tag_reports_offending_paragraph(self):
        document = Document()
        document.add_paragraph("Earlier {{ tag }}")
        document.add_paragraph("{%p fro x in y %}")
        template = io.BytesIO()
        document.save(template)

        with self.assertRaises(ValueError) as raised:
            render_docx(template.getvalue(), {})

        self.assertIn('template text "{%p fro x in y %}"', str(raised.exception))

    def test_docx_unknown_tag_prefers_later_offending_paragraph(self):
        document = Document()
        for name in ("first", "second"):
            document.add_paragraph("{%% for %s in values %%}" % name)
            document.add_paragraph("{{ %s }}" % name)
            document.add_paragraph("{% endfor %}")
        document.add_paragraph("{% frobnicate x %}")
        template = io.BytesIO()
        document.save(template)

        with self.assertRaises(ValueError) as raised:
            render_docx(template.getvalue(), {})

        self.assertIn(
            'template text "{% frobnicate x %}"',
            str(raised.exception),
        )


    def test_docx_runtime_error_uses_runtime_prefix(self):
        document = Document()
        document.add_paragraph("{{ 1 // 0 }}")
        template = io.BytesIO()
        document.save(template)

        with self.assertRaises(ValueError) as raised:
            render_docx(template.getvalue(), {})

        message = str(raised.exception)
        self.assertTrue(
            message.startswith(
                "DOCX template could not be rendered with this specification:"
            ),
            message,
        )
        self.assertIn("{{ 1 // 0 }}", message)
        self.assertIn("integer division or modulo by zero", message)

    def test_clean_docx_renders(self):
        document = Document()
        document.add_paragraph("Hello {{ name }}")
        template = io.BytesIO()
        document.save(template)

        output = render_docx(template.getvalue(), {"name": "world"})

        self.assertTrue(output.startswith(b"PK"))
        rendered = Document(io.BytesIO(output))
        self.assertEqual(rendered.paragraphs[0].text, "Hello world")

    def test_xlsx_broken_cell_reports_sheet_and_coordinate(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Summary"
        sheet["B3"] = "{{ 1 // 0 }}"
        template = io.BytesIO()
        workbook.save(template)

        with self.assertRaises(ValueError) as raised:
            render_xlsx(template.getvalue(), {})

        message = str(raised.exception)
        self.assertTrue(
            message.startswith(
                "XLSX template could not be rendered with this specification:"
            ),
            message,
        )
        self.assertIn("Summary", message)
        self.assertIn("B3", message)
        self.assertIn("{{ 1 // 0 }}", message)
        self.assertIn("integer division or modulo by zero", message)

    def test_clean_xlsx_renders(self):
        workbook = Workbook()
        workbook.active["A1"] = "Hello {{ name }}"
        template = io.BytesIO()
        workbook.save(template)

        output = render_xlsx(template.getvalue(), {"name": "world"})

        self.assertTrue(output.startswith(b"PK"))
        rendered = load_workbook(io.BytesIO(output))
        self.assertEqual(rendered.active["A1"].value, "Hello world")
